from pathlib import Path
from io import BytesIO

IMAGES = {".jpg", ".jpeg", ".png", ".heic", ".heif", ".webp"}
VIDEOS = {".mp4", ".mov", ".m4v", ".mkv", ".webm", ".avi"}
AUDIO = {".mp3", ".wav", ".m4a", ".flac", ".ogg", ".aac", ".aiff", ".aif", ".opus"}
TEXT = {".txt", ".md", ".markdown", ".rst", ".json", ".yaml", ".yml", ".toml",
        ".py", ".swift", ".js", ".jsx", ".ts", ".tsx", ".sh", ".zsh", ".go", ".rs",
        ".c", ".h", ".cpp", ".css", ".html", ".sql", ".ini", ".conf"}
EXCLUDED = {"node_modules", ".git", ".venv", "venv", "__pycache__", "dist", "build",
            ".next", ".cache", "vendor", "Pods", "DerivedData"}


def kind(path: Path):
    suffix = path.suffix.lower()
    if suffix in IMAGES:
        return "image"
    if suffix in VIDEOS:
        return "video"
    if suffix in AUDIO:
        return "audio"
    if suffix in TEXT or suffix in {".pdf", ".docx"}:
        return "document"
    return None


def extract(path: Path):
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if path.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        parts.extend(" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
        return "\n".join(parts)
    return path.read_text(encoding="utf-8-sig")


def thumbnail(path: Path, size=256, quality=60):
    from PIL import Image, ImageOps
    from .images import open_image
    with open_image(path) as image:
        image = ImageOps.exif_transpose(image).convert("RGB")
        image.thumbnail((size, size))
        output = BytesIO()
        image.save(output, format="JPEG", quality=quality, optimize=True)
        return output.getvalue()
